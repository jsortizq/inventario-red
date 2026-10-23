#!/usr/bin/env python3
"""
Inventario de red
=================

Recorre un rango de red, encuentra los equipos encendidos y para cada uno
registra: IP, nombre (DNS inverso), dirección MAC, fabricante aproximado,
puertos abiertos y el tipo de equipo probable. El resultado queda en un
Excel con dos hojas (equipos y resumen) y, si se pide, en CSV.

Lo uso para levantar el inventario cuando llego a una red nueva o para
comparar contra el inventario anterior y ver qué equipos aparecieron.

Uso:
    python inventario_red.py 192.168.1.0/24
    python inventario_red.py 192.168.1.0/24 10.0.0.0/28 -o inventario.xlsx --csv
    python inventario_red.py 192.168.1.10-192.168.1.60 --hilos 128

Autor: Juan Sebastián Ortiz Quintero
"""
from __future__ import annotations

import argparse
import csv
import ipaddress
import platform
import re
import shutil
import socket
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime

# Puertos que dicen mucho de un equipo en una red de oficina o planta
PUERTOS = {
    21: "FTP",
    22: "SSH",
    23: "Telnet",
    80: "HTTP",
    102: "Siemens S7",
    135: "RPC",
    139: "NetBIOS",
    443: "HTTPS",
    445: "SMB",
    502: "Modbus TCP",
    554: "RTSP (cámara)",
    1433: "SQL Server",
    3306: "MySQL",
    3389: "Escritorio remoto",
    8000: "Hikvision / HTTP",
    8080: "HTTP alterno",
    9100: "Impresora (RAW)",
    37777: "Dahua",
    44818: "EtherNet/IP (Allen-Bradley)",
}

# Prefijos MAC (OUI) frecuentes. Lista corta a propósito: cubre lo que más se ve en pymes.
FABRICANTES = {
    "00:1B:21": "Intel", "3C:52:82": "HP", "00:50:56": "VMware", "00:0C:29": "VMware",
    "08:00:27": "VirtualBox", "B8:27:EB": "Raspberry Pi", "DC:A6:32": "Raspberry Pi",
    "00:1C:06": "Siemens", "00:0E:8C": "Siemens", "00:00:BC": "Rockwell Automation",
    "00:80:F4": "Schneider Electric", "BC:AD:28": "Hikvision", "44:19:B6": "Hikvision",
    "3C:EF:8C": "Dahua", "E0:50:8B": "Dahua", "F4:F2:6D": "TP-Link", "50:C7:BF": "TP-Link",
    "24:5A:4C": "Ubiquiti", "FC:EC:DA": "Ubiquiti", "4C:5E:0C": "MikroTik", "00:0C:42": "MikroTik",
    "00:1E:0B": "HP (impresora)", "00:26:AB": "Epson", "00:00:48": "Epson", "00:80:77": "Brother",
    "D4:BE:D9": "Dell", "F8:BC:12": "Dell", "00:15:5D": "Hyper-V", "70:85:C2": "ASRock",
}

ES_WINDOWS = platform.system().lower() == "windows"


@dataclass
class Equipo:
    ip: str
    nombre: str = ""
    mac: str = ""
    fabricante: str = ""
    latencia_ms: float | None = None
    puertos: list[int] = field(default_factory=list)

    @property
    def servicios(self) -> str:
        return ", ".join(f"{p} {PUERTOS[p]}" for p in self.puertos)

    @property
    def tipo(self) -> str:
        """Clasificación aproximada a partir de puertos y fabricante."""
        p = set(self.puertos)
        f = self.fabricante.lower()
        if p & {102, 502, 44818} or any(x in f for x in ("siemens", "rockwell", "schneider")):
            return "PLC / equipo industrial"
        if p & {554, 8000, 37777} or any(x in f for x in ("hikvision", "dahua")):
            return "Cámara / NVR"
        if 9100 in p or any(x in f for x in ("epson", "brother", "impresora")):
            return "Impresora"
        if 1433 in p or 3306 in p:
            return "Servidor de base de datos"
        if 445 in p and 3389 in p:
            return "Servidor / PC Windows"
        if 445 in p or 135 in p:
            return "PC Windows"
        if any(x in f for x in ("ubiquiti", "mikrotik", "tp-link")):
            return "Red (router / AP / switch)"
        if 22 in p and not p & {445, 135}:
            return "Linux / equipo de red"
        if p & {80, 443}:
            return "Equipo con interfaz web"
        return "Desconocido"


# ---------------------------------------------------------------- descubrimiento

HAY_PING = shutil.which("ping") is not None


def sondeo_tcp(ip: str, timeout_ms: int) -> float | None:
    """Alternativa al ping (sin el comando o con ICMP bloqueado): intenta puertos comunes."""
    inicio = time.perf_counter()
    for p in (445, 80, 443, 22, 135, 3389, 554):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(timeout_ms / 1000 / 3)
            if s.connect_ex((ip, p)) in (0, 111, 10061):   # abierto, o rechazado: el equipo existe
                return round((time.perf_counter() - inicio) * 1000, 1)
    return None


def ping(ip: str, timeout_ms: int) -> float | None:
    """Devuelve la latencia en ms o None si no responde."""
    if not HAY_PING:
        return sondeo_tcp(ip, timeout_ms)
    if ES_WINDOWS:
        cmd = ["ping", "-n", "1", "-w", str(timeout_ms), ip]
    else:
        cmd = ["ping", "-c", "1", "-W", str(max(1, timeout_ms // 1000)), ip]
    inicio = time.perf_counter()
    try:
        r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                           text=True, timeout=timeout_ms / 1000 + 2)
    except (subprocess.TimeoutExpired, FileNotFoundError):
        return None
    if r.returncode != 0 or "TTL=" not in r.stdout.upper():
        return None
    m = re.search(r"(?:time|tiempo)[=<]\s*([\d.,]+)\s*ms", r.stdout, re.IGNORECASE)
    if m:
        return float(m.group(1).replace(",", "."))
    return round((time.perf_counter() - inicio) * 1000, 1)


def puerto_abierto(ip: str, puerto: int, timeout: float) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(timeout)
        return s.connect_ex((ip, puerto)) == 0


def nombre_dns(ip: str) -> str:
    try:
        return socket.gethostbyaddr(ip)[0]
    except (socket.herror, socket.gaierror, OSError):
        return ""


def tabla_arp() -> dict[str, str]:
    """Lee la tabla ARP del sistema: {ip: mac}. Funciona en Windows y Linux."""
    macs: dict[str, str] = {}
    try:
        if ES_WINDOWS:
            salida = subprocess.run(["arp", "-a"], capture_output=True, text=True).stdout
        else:
            salida = subprocess.run(["ip", "neigh"], capture_output=True, text=True).stdout
    except FileNotFoundError:
        return macs
    patron = re.compile(r"(\d+\.\d+\.\d+\.\d+)\D+?(([0-9a-fA-F]{2}[:-]){5}[0-9a-fA-F]{2})")
    for linea in salida.splitlines():
        m = patron.search(linea)
        if m:
            macs[m.group(1)] = m.group(2).replace("-", ":").upper()
    return macs


def fabricante(mac: str) -> str:
    return FABRICANTES.get(mac[:8], "") if mac else ""


# ---------------------------------------------------------------- rangos

def expandir_rangos(rangos: list[str]) -> list[str]:
    """Acepta CIDR (192.168.1.0/24), rango (192.168.1.10-192.168.1.50) o IP suelta."""
    ips: list[str] = []
    for r in rangos:
        if "-" in r:
            a, b = (ipaddress.IPv4Address(x.strip()) for x in r.split("-", 1))
            if b < a:
                a, b = b, a
            ips += [str(ipaddress.IPv4Address(i)) for i in range(int(a), int(b) + 1)]
        elif "/" in r:
            red = ipaddress.IPv4Network(r, strict=False)
            hosts = list(red.hosts()) or [red.network_address]
            ips += [str(h) for h in hosts]
        else:
            ips.append(str(ipaddress.IPv4Address(r)))
    return list(dict.fromkeys(ips))  # sin duplicados, en orden


# ---------------------------------------------------------------- escaneo

def escanear(ips: list[str], hilos: int, timeout_ms: int, puertos: list[int]) -> list[Equipo]:
    vivos: list[Equipo] = []
    total = len(ips)

    print(f"Buscando equipos en {total} direcciones...")
    with ThreadPoolExecutor(max_workers=hilos) as ex:
        futuros = {ex.submit(ping, ip, timeout_ms): ip for ip in ips}
        for n, fut in enumerate(as_completed(futuros), 1):
            lat = fut.result()
            if lat is not None:
                vivos.append(Equipo(ip=futuros[fut], latencia_ms=lat))
            if n % 25 == 0 or n == total:
                print(f"\r  {n}/{total} revisadas, {len(vivos)} encendidas", end="", flush=True)
    print()

    if not vivos:
        return vivos

    print(f"Revisando puertos y nombres de {len(vivos)} equipos...")
    timeout_s = max(0.3, timeout_ms / 1000 / 2)

    def detallar(eq: Equipo) -> Equipo:
        eq.nombre = nombre_dns(eq.ip)
        eq.puertos = [p for p in puertos if puerto_abierto(eq.ip, p, timeout_s)]
        return eq

    with ThreadPoolExecutor(max_workers=min(hilos, 32)) as ex:
        list(ex.map(detallar, vivos))

    arp = tabla_arp()
    for eq in vivos:
        eq.mac = arp.get(eq.ip, "")
        eq.fabricante = fabricante(eq.mac)

    vivos.sort(key=lambda e: ipaddress.IPv4Address(e.ip))
    return vivos


# ---------------------------------------------------------------- exportar

COLUMNAS = ["IP", "Nombre", "MAC", "Fabricante", "Tipo probable", "Latencia (ms)", "Puertos abiertos"]


def filas(equipos: list[Equipo]):
    for e in equipos:
        yield [e.ip, e.nombre, e.mac, e.fabricante, e.tipo, e.latencia_ms, e.servicios]


def exportar_csv(equipos: list[Equipo], ruta: str) -> None:
    with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUMNAS)
        w.writerows(filas(equipos))


def exportar_excel(equipos: list[Equipo], ruta: str, rangos: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Equipos"
    ws.append(COLUMNAS)
    for fila in filas(equipos):
        ws.append(fila)

    if equipos:
        tabla = Table(displayName="Equipos", ref=f"A1:{get_column_letter(len(COLUMNAS))}{len(equipos) + 1}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabla)

    anchos = [16, 30, 19, 20, 26, 13, 60]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"

    # Resaltar servicios inseguros expuestos
    alerta = PatternFill("solid", fgColor="FDE2E1")
    for fila in ws.iter_rows(min_row=2):
        texto = str(fila[6].value or "")
        if "Telnet" in texto or "FTP" in texto:
            fila[6].fill = alerta

    # Hoja de resumen
    rs = wb.create_sheet("Resumen")
    rs["A1"] = "Inventario de red"
    rs["A1"].font = Font(size=14, bold=True)
    rs.append([])
    rs.append(["Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")])
    rs.append(["Rangos", ", ".join(rangos)])
    rs.append(["Equipos encontrados", len(equipos)])
    rs.append([])
    rs.append(["Tipo probable", "Cantidad"])
    rs.cell(rs.max_row, 1).font = rs.cell(rs.max_row, 2).font = Font(bold=True)
    conteo: dict[str, int] = {}
    for e in equipos:
        conteo[e.tipo] = conteo.get(e.tipo, 0) + 1
    for tipo, n in sorted(conteo.items(), key=lambda x: -x[1]):
        rs.append([tipo, n])
    inseguros = [e.ip for e in equipos if {21, 23} & set(e.puertos)]
    if inseguros:
        rs.append([])
        rs.append(["Equipos con FTP/Telnet abiertos (revisar)", ", ".join(inseguros)])
        rs.cell(rs.max_row, 1).font = Font(bold=True, color="B3261E")
    rs.column_dimensions["A"].width = 42
    rs.column_dimensions["B"].width = 50
    for fila in rs.iter_rows():
        for c in fila:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(ruta)


# ---------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description="Inventario de equipos de una red local.")
    ap.add_argument("rangos", nargs="+", help="CIDR, rango inicio-fin o IP. Ej: 192.168.1.0/24")
    ap.add_argument("-o", "--salida", default=f"inventario_{datetime.now():%Y%m%d_%H%M}.xlsx",
                    help="Archivo Excel de salida")
    ap.add_argument("--csv", action="store_true", help="Guardar también en CSV")
    ap.add_argument("--hilos", type=int, default=64, help="Pings simultáneos (por defecto 64)")
    ap.add_argument("--timeout", type=int, default=800, help="Tiempo de espera en ms (por defecto 800)")
    ap.add_argument("--puertos", help="Lista de puertos separada por comas (por defecto, la lista interna)")
    args = ap.parse_args()

    try:
        ips = expandir_rangos(args.rangos)
    except ValueError as e:
        print(f"Rango inválido: {e}", file=sys.stderr)
        return 2
    if len(ips) > 4096:
        print(f"El rango tiene {len(ips)} direcciones; divídalo en partes de /20 o menos.", file=sys.stderr)
        return 2

    puertos = sorted(int(p) for p in args.puertos.split(",")) if args.puertos else sorted(PUERTOS)
    for p in puertos:
        PUERTOS.setdefault(p, "")

    inicio = time.perf_counter()
    equipos = escanear(ips, args.hilos, args.timeout, puertos)
    duracion = time.perf_counter() - inicio

    print(f"\n{'IP':<16}{'Nombre':<28}{'Tipo probable':<28}Puertos")
    for e in equipos:
        print(f"{e.ip:<16}{(e.nombre or '-')[:27]:<28}{e.tipo:<28}{','.join(map(str, e.puertos)) or '-'}")

    try:
        exportar_excel(equipos, args.salida, args.rangos)
        print(f"\nExcel: {args.salida}")
    except ImportError:
        print("\nopenpyxl no está instalado (pip install openpyxl); se guarda solo CSV.")
        args.csv = True
    if args.csv:
        ruta_csv = args.salida.rsplit(".", 1)[0] + ".csv"
        exportar_csv(equipos, ruta_csv)
        print(f"CSV:   {ruta_csv}")

    print(f"{len(equipos)} equipos en {duracion:.1f} s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
